"""Build the QueryGuard SRS presentation (8 slides, 4 per presenter).

Design: editorial dark, the settled deck look (bg 121418, panels 1B2027,
ink F2F0E8, accent E74C3C, slate 9FB3C8). Labeled-block slide style.
Rerun after any edit:  python presentation/build_deck.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "QueryGuard_SRS.pptx"

BG = RGBColor(0x12, 0x14, 0x18)
PANEL = RGBColor(0x1B, 0x20, 0x27)
INK = RGBColor(0xF2, 0xF0, 0xE8)
SUB = RGBColor(0xA8, 0xB2, 0xBB)
SLATE = RGBColor(0x9F, 0xB3, 0xC8)
RED = RGBColor(0xE7, 0x4C, 0x3C)
GREEN = RGBColor(0x43, 0xD1, 0x7C)
AMBER = RGBColor(0xE8, 0xB3, 0x41)
ORANGE = RGBColor(0xE5, 0x91, 0x3C)

W, H = Inches(13.333), Inches(7.5)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def box(slide, l, t, w, h, text, size, color=INK, bold=False, font="Segoe UI",
        align=PP_ALIGN.LEFT, caps=False, spacing=None):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = line.upper() if caps else line
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        f.name = font
    return tb


def label(slide, text, color=SLATE):
    box(slide, Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.4),
        text, 13, color, bold=True, caps=True)


def panel(slide, l, t, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = PANEL
    sh.line.color.rgb = RGBColor(0x2A, 0x31, 0x3A)
    sh.shadow.inherit = False
    return sh


def block(slide, l, t, w, h, tag, body, tag_color=SLATE):
    panel(slide, l, t, w, h)
    box(slide, l + Inches(0.25), t + Inches(0.18), w - Inches(0.5), Inches(0.35),
        tag, 12, tag_color, bold=True, caps=True)
    box(slide, l + Inches(0.25), t + Inches(0.62), w - Inches(0.5), h - Inches(0.75),
        body, 14, INK)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H

# S1 · title (Arwa)
s = new_slide(prs)
box(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.4),
    "CS5391 · Group Project 1 · Requirement Document", 14, SLATE, bold=True, caps=True)
box(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.2),
    "QueryGuard", 54, ORANGE, bold=True, font="Georgia")
box(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(0.9),
    "An AI writes SQL from your question. A safety gate checks it\nbefore anything touches the database.",
    20, INK)
box(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.8),
    "Mahd Amjad  ·  Arwa Ali Arafeh\nCS5391 Survey of Software Engineering · Fall 2026 · Dr. Xiao Chen",
    15, SUB)

# S2 · the problem (Arwa)
s = new_slide(prs)
label(s, "The problem · Arwa")
box(s, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.7),
    "Letting people ask a database questions in plain English", 24, INK, bold=True)
block(s, Inches(0.7), Inches(1.9), Inches(3.83), Inches(2.1),
      "Useful", "Anyone can get answers without knowing SQL. The AI writes the query for them.")
block(s, Inches(4.75), Inches(1.9), Inches(3.83), Inches(2.1),
      "Dangerous", "One wrong or malicious query deletes records, or walks SSNs out the front door.", RED)
block(s, Inches(8.8), Inches(1.9), Inches(3.83), Inches(2.1),
      "The gap", "Nothing in that pipeline checks the AI's work. Unsafe by default.", AMBER)
box(s, Inches(0.7), Inches(4.4), Inches(11.9), Inches(0.9),
    "So: keep the AI, add a gate.\nA deterministic check between the AI and the database.", 18, SLATE)

# S3 · what QueryGuard does (Arwa)
s = new_slide(prs)
label(s, "What QueryGuard does · Arwa")
box(s, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.7),
    "Five steps, one recorded path", 24, INK, bold=True)
steps = ["1. You ask a question", "2. The AI drafts the SQL", "3. The gate checks it",
         "4. Safe queries run, read-only", "5. Everything is logged"]
for i, step in enumerate(steps):
    x = Inches(0.7 + i * 2.45)
    panel(s, x, Inches(2.0), Inches(2.25), Inches(1.0))
    box(s, x + Inches(0.15), Inches(2.2), Inches(1.95), Inches(0.7), step, 13, INK)
box(s, Inches(0.7), Inches(3.4), Inches(11.9), Inches(0.5),
    "Three outcomes, in plain words:", 18, INK, bold=True)
outcomes = [("PASS", "ran as asked", GREEN), ("MASK", "answered, private columns removed", AMBER),
            ("BLOCK", "stopped, nothing touched", RED)]
for i, (v, meaning, c) in enumerate(outcomes):
    y = Inches(3.95 + i * 0.62)
    panel(s, Inches(0.7), y, Inches(1.7), Inches(0.5))
    box(s, Inches(0.85), y + Inches(0.08), Inches(1.4), Inches(0.35), v, 14, c, bold=True)
    box(s, Inches(2.6), y + Inches(0.08), Inches(9.5), Inches(0.35), meaning, 14, INK)

# S4 · user stories (Arwa)
s = new_slide(prs)
label(s, "User stories · Arwa")
box(s, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.7),
    "Who wants what, and why", 24, INK, bold=True)
stories = [
    ("A demo viewer", "to ask a question and see an answer", "the product works end to end"),
    ("A privacy-conscious viewer", "to know SSNs can never appear in answers", "the privacy claim is believed"),
    ("A compliance reviewer", "to browse a log of every decision", "what ran and why is verifiable later"),
    ("An evaluator", "to run one command against known attacks", "safety is measured, not asserted"),
]
for i, (who, want, why) in enumerate(stories):
    y = Inches(1.9 + i * 1.05)
    panel(s, Inches(0.7), y, Inches(11.9), Inches(0.9))
    box(s, Inches(0.95), y + Inches(0.12), Inches(3.2), Inches(0.6), "As " + who, 14, SLATE, bold=True)
    box(s, Inches(4.3), y + Inches(0.12), Inches(4.6), Inches(0.6), "I want " + want, 14, INK)
    box(s, Inches(9.1), y + Inches(0.12), Inches(3.3), Inches(0.6), "so that " + why, 13, SUB)

# S5 · the nine rules (Mahd)
s = new_slide(prs)
label(s, "The gate · Mahd")
box(s, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.7),
    "Nine rules, checked on every query", 24, INK, bold=True)
rules = [
    ("No destructive statements", "DELETE, UPDATE, DROP never run. Read-only database"),
    ("One query at a time", "no stacked commands"),
    ("No hidden write commands", "comment tricks are caught by re-parsing"),
    ("Only allowed tables", "everything outside the clinic tables is refused"),
    ("Private columns stay private", "masked, and indirect routes (grouping, ordering) blocked"),
    ("No system-table snooping", "no mapping the database structure"),
    ("No execution tricks", "no code execution, file attaches, or settings"),
    ("Huge answers get a hint", "unbounded reads get a LIMIT suggestion"),
    ("Unreadable means blocked", "if the gate cannot parse it, nothing runs"),
]
for i, (name, what) in enumerate(rules):
    y = Inches(1.8 + i * 0.55)
    box(s, Inches(0.7), y, Inches(4.3), Inches(0.45), name, 14, SLATE, bold=True)
    box(s, Inches(5.1), y, Inches(7.4), Inches(0.45), what, 14, INK)

# S6 · how it is built (Mahd)
s = new_slide(prs)
label(s, "How it is built · Mahd")
box(s, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.7),
    "Boring on purpose, verified everywhere", 24, INK, bold=True)
block(s, Inches(0.7), Inches(1.9), Inches(5.8), Inches(1.6),
      "Rules on the parsed query", "Rules run on a parsed SQL tree, never on raw text. Comment tricks and hidden writes survive no other way.")
block(s, Inches(6.8), Inches(1.9), Inches(5.8), Inches(1.6),
      "Read-only, one path", "The executor holds a read-only connection. Blocked queries never reach the database. Every decision is logged first.")
block(s, Inches(0.7), Inches(3.7), Inches(5.8), Inches(1.6),
      "No external dependencies", "A deterministic stand-in stands in for the live AI, so tests and demos never depend on an API.")
block(s, Inches(6.8), Inches(3.7), Inches(5.8), Inches(1.6),
      "Continuous integration", "Every push runs 40 tests plus a 64-case labeled corpus. Any miss fails the build.")
box(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.6),
    "Practices from this course: user stories, test-first development, refactoring, continuous integration.", 15, SLATE)

# S7 · measured results (Mahd)
s = new_slide(prs)
label(s, "Measured results · Mahd")
box(s, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.7),
    "Measured, not asserted", 24, INK, bold=True)
stats = [("34 / 34", "attacks stopped", GREEN), ("0 / 30", "safe questions blocked", GREEN),
         ("0.57 ms", "slowest check (target 50)", SLATE), ("40", "automated tests", INK),
         ("384", "rewrite-stability checks", INK)]
for i, (n, l, c) in enumerate(stats):
    x = Inches(0.7 + i * 2.45)
    panel(s, x, Inches(1.9), Inches(2.25), Inches(1.5))
    box(s, x + Inches(0.15), Inches(2.05), Inches(1.95), Inches(0.6), n, 26, c, bold=True)
    box(s, x + Inches(0.15), Inches(2.75), Inches(1.95), Inches(0.5), l, 12, SUB)
box(s, Inches(0.7), Inches(3.8), Inches(11.9), Inches(1.0),
    "One honest note: 9 of the 34 attacks were answered with private columns removed (MASK)\ninstead of blocked. Nothing leaked. That is the designed outcome, reported separately so the numbers stay honest.",
    15, SUB)
box(s, Inches(0.7), Inches(5.0), Inches(11.9), Inches(0.8),
    "The live AI mode is not built yet. The demo runs on a deterministic stand-in,\nso nothing on this slide depends on an external service.", 15, SLATE)

# S8 · close (Mahd)
s = new_slide(prs)
label(s, "What is next · Mahd")
box(s, Inches(0.7), Inches(1.1), Inches(11.9), Inches(0.8),
    "The demo is the argument.", 28, INK, bold=True)
nexts = [("Live AI mode", "swap the stand-in for a real model behind the same gate"),
         ("Intent matching", "does the SQL do what the question asked (stretch)"),
         ("Corpus hardening", "fuzzed attack variants, so the guard tests itself")]
for i, (t2, d) in enumerate(nexts):
    y = Inches(2.3 + i * 0.95)
    panel(s, Inches(0.7), y, Inches(11.9), Inches(0.8))
    box(s, Inches(0.95), y + Inches(0.16), Inches(3.4), Inches(0.5), t2, 15, SLATE, bold=True)
    box(s, Inches(4.5), y + Inches(0.16), Inches(7.9), Inches(0.5), d, 14, INK)
box(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.9),
    "Live demo: safe question, blocked attack, masked private columns, the log.\nMahd Amjad · Arwa Ali Arafeh · github.com/Mahd-Amjad/queryguard-cs5391",
    15, SUB)

prs.save(OUT)
print(f"built {OUT} with {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
